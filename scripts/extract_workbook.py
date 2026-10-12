#!/usr/bin/env python3
"""Regenerate use-cases.json from the GNOSIS demo workbook.

The website reads the JSON file. It does not open the Excel file at runtime.
Install the one dependency, then run this from the repository root:

    python3 -m pip install openpyxl
    python3 scripts/extract_workbook.py --input "/path/to/GNOSIS Demo Use Cases.xlsx"
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.stderr.write(
        "openpyxl is required. Install it with: python3 -m pip install openpyxl\n"
    )
    raise SystemExit(1)


VIDEO_URL = "https://www.youtube.com/watch?v=sHaHzk2zHg4&t=1s"


def as_text(value):
    if value is None:
        return ""
    return str(value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Path to the Excel workbook")
    parser.add_argument(
        "--output",
        default="use-cases.json",
        help="Where to write the JSON (default: use-cases.json)",
    )
    args = parser.parse_args()

    source = Path(args.input)
    if not source.is_file():
        sys.stderr.write("Workbook not found: %s\n" % source)
        return 1

    workbook = openpyxl.load_workbook(source, data_only=False, read_only=False)
    if "Demo" not in workbook.sheetnames:
        sys.stderr.write("Expected a sheet named Demo. Found: %s\n" % workbook.sheetnames)
        return 1

    sheet = workbook["Demo"]
    headers = [as_text(sheet.cell(5, col).value).strip() for col in range(1, 6)]
    expected_headers = [
        "Use Case ID",
        "Use Case",
        "Question",
        "GNOSIS Demo Response",
        "Gemini (3.6 Flash) Response",
    ]
    if headers != expected_headers:
        sys.stderr.write("Unexpected column headers: %r\n" % headers)
        return 1

    use_cases = []
    for row in range(6, 18):
        raw_id = sheet.cell(row, 1).value
        try:
            case_id = int(raw_id)
        except (TypeError, ValueError):
            sys.stderr.write("Row %s has an invalid use case id: %r\n" % (row, raw_id))
            return 1
        case = {
            "id": case_id,
            "title": as_text(sheet.cell(row, 2).value),
            "question": as_text(sheet.cell(row, 3).value),
            "gnosis": as_text(sheet.cell(row, 4).value),
            "gemini": as_text(sheet.cell(row, 5).value),
        }
        missing = [key for key in ("title", "question", "gnosis", "gemini") if not case[key].strip()]
        if missing:
            sys.stderr.write("Row %s is missing: %s\n" % (row, ", ".join(missing)))
            return 1
        use_cases.append(case)

    if [case["id"] for case in use_cases] != list(range(1, 13)):
        sys.stderr.write("Expected use case ids 1 through 12.\n")
        return 1

    payload = {
        "source": {
            "title": as_text(sheet["A1"].value).strip(),
            "context": as_text(sheet["A2"].value).strip(),
            "bibliography": [as_text(sheet["A3"].value).strip()],
            "columns": {
                "gnosis": headers[3],
                "gemini": headers[4],
            },
        },
        "models": {
            "gnosis": {
                "name": "GNOSIS",
                "responseLabel": headers[3],
                "role": "Scientific demonstration system",
            },
            "gemini": {
                "name": "Gemini",
                "responseLabel": headers[4],
                "role": "General-purpose model",
            },
        },
        "collaborationUrl": "",
        "videoUrl": VIDEO_URL,
        "useCases": use_cases,
    }

    output = Path(args.output)
    output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print("Wrote %s" % output)
    print("Context: %s" % payload["source"]["context"])
    print("Bibliography: %s" % payload["source"]["bibliography"][0])
    for case in use_cases:
        print(
            "%2d  %-48s  Q:%4d  GNOSIS:%5d  Gemini:%5d"
            % (
                case["id"],
                case["title"][:48],
                len(case["question"]),
                len(case["gnosis"]),
                len(case["gemini"]),
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
